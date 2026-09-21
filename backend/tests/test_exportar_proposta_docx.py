"""
Teste do endpoint GET /api/editais/{id}/proposta.docx — chama a função da
rota direto (sem HTTP, mesmo padrão de test_exportar_proposta_pdf.py).
Rode com:  cd backend && pytest
"""
import asyncio
import io

from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.main import exportar_proposta_docx
from app.models import Base, Usuario, Edital, ItemEdital, Proposta


async def _drenar(body_iterator):
    partes = []
    async for pedaco in body_iterator:
        partes.append(pedaco)
    return b"".join(partes)


def _sessao():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    return Session()


def test_exportar_proposta_docx_retorna_docx_valido():
    db = _sessao()
    u = Usuario(nome="Empresa Teste", email="e@t.com", senha_hash="x")
    db.add(u)
    db.commit()
    ed = Edital(fonte="PNCP", id_externo="ed1", orgao="Orgao Teste",
               objeto="Aquisicao", uf="SP")
    db.add(ed)
    db.commit()
    db.add(ItemEdital(edital_id=ed.id, numero=1, descricao="Papel A4",
                      quantidade=10, valor_unitario=25.0))
    db.commit()

    resp = exportar_proposta_docx(ed.id, user=u, db=db)

    assert resp.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    corpo = asyncio.run(_drenar(resp.body_iterator))
    assert corpo[:2] == b"PK"
    Document(io.BytesIO(corpo))   # não pode lançar exceção ao reabrir


def test_docx_nao_pode_ficar_em_cache_no_navegador():
    """Mesmo achado real do PDF (test_pdf_nao_pode_ficar_em_cache_no_
    navegador) -- proposta é gerada com dado ao vivo do catálogo a cada
    request, nunca pode ficar em cache."""
    db = _sessao()
    u = Usuario(nome="Empresa Teste", email="e5@t.com", senha_hash="x")
    db.add(u)
    db.commit()
    ed = Edital(fonte="PNCP", id_externo="ed5", orgao="Orgao Teste",
               objeto="Aquisicao", uf="SP")
    db.add(ed)
    db.commit()
    db.add(ItemEdital(edital_id=ed.id, numero=1, descricao="Papel A4",
                      quantidade=10, valor_unitario=25.0))
    db.commit()

    resp = exportar_proposta_docx(ed.id, user=u, db=db)
    assert resp.headers.get("cache-control") == "no-store"


def _texto_do_docx(corpo: bytes) -> str:
    doc = Document(io.BytesIO(corpo))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                partes.append(celula.text)
    return "\n".join(partes)


def test_docx_mostra_numero_do_item_na_tabela():
    db = _sessao()
    u = Usuario(nome="Empresa Teste", email="e2@t.com", senha_hash="x")
    db.add(u)
    db.commit()
    ed = Edital(fonte="PNCP", id_externo="ed2", orgao="Orgao Teste",
               objeto="Aquisicao", uf="SP")
    db.add(ed)
    db.commit()
    db.add(ItemEdital(edital_id=ed.id, numero=24, descricao="Papel A4",
                      quantidade=10, valor_unitario=25.0))
    db.commit()

    resp = exportar_proposta_docx(ed.id, user=u, db=db)
    corpo = asyncio.run(_drenar(resp.body_iterator))
    texto = _texto_do_docx(corpo)
    assert "24" in texto


def test_docx_nao_quebra_com_item_sem_numero():
    db = _sessao()
    u = Usuario(nome="Empresa Teste", email="e3@t.com", senha_hash="x")
    db.add(u)
    db.commit()
    ed = Edital(fonte="PNCP", id_externo="ed3", orgao="Orgao Teste",
               objeto="Aquisicao", uf="SP")
    db.add(ed)
    db.commit()
    db.add(Proposta(edital_id=ed.id, usuario_id=u.id, itens=[
        {"descricao": "Item digitado à mão, sem número", "quantidade": 1,
         "custo_unit": 0, "preco_unit": 10.0},
    ]))
    db.commit()

    resp = exportar_proposta_docx(ed.id, user=u, db=db)   # não pode lançar exceção

    corpo = asyncio.run(_drenar(resp.body_iterator))
    assert corpo[:2] == b"PK"


def test_docx_edital_inexistente_retorna_404():
    from fastapi import HTTPException
    import pytest as _pytest
    db = _sessao()
    u = Usuario(nome="Empresa Teste", email="e6@t.com", senha_hash="x")
    db.add(u)
    db.commit()

    with _pytest.raises(HTTPException) as exc:
        exportar_proposta_docx(999999, user=u, db=db)
    assert exc.value.status_code == 404

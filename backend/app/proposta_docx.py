"""
Gera a proposta comercial em DOCX (Word), mesmo conteúdo/ordem de
gerar_pdf_proposta (proposta_pdf.py) -- pedido do usuário: algumas
licitações pedem a proposta em formato editável, pra imprimir em papel
timbrado próprio ou ajustar um detalhe de última hora sem depender de
gerar o PDF de novo aqui. Reaproveita os helpers puros (_fmt_moeda,
_abreviar_unidade, _MESES, _DECLARACAO) de proposta_pdf.py -- só a
montagem do documento em si é específica de cada formato.
"""
from __future__ import annotations
import base64
import datetime
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .proposta_pdf import _DECLARACAO, _MESES, _abreviar_unidade, _fmt_moeda

_COR_MUTED = RGBColor(0x5B, 0x67, 0x70)
_COR_TXT = RGBColor(0x14, 0x19, 0x1E)


def _decodificar_logo(data_uri: str | None) -> io.BytesIO | None:
    """Mesma regra de proposta_pdf._decodificar_logo -- None se vazio,
    malformado, ou SVG (python-docx também não abre SVG arbitrário)."""
    if not data_uri or not data_uri.startswith("data:image/"):
        return None
    m = re.match(r"data:image/(\w+);base64,(.+)$", data_uri, re.DOTALL)
    if not m:
        return None
    tipo, b64 = m.group(1).lower(), m.group(2)
    if tipo in ("svg+xml", "svg"):
        return None
    try:
        return io.BytesIO(base64.b64decode(b64))
    except Exception:
        return None


def _texto(paragrafo, texto: str, *, negrito=False, tamanho=10, cor=None):
    run = paragrafo.add_run(texto)
    run.bold = negrito
    run.font.size = Pt(tamanho)
    if cor is not None:
        run.font.color.rgb = cor
    return run


def _celula_sombreada(celula, hex_cor="F0F2F5"):
    """Preenchimento de fundo da célula -- python-docx não expõe isso por
    cima, precisa mexer direto no XML da célula (mesmo padrão usado pelo
    cabeçalho em negrito da tabela do PDF)."""
    tcPr = celula._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_cor})
    tcPr.append(shd)


def gerar_docx_proposta(remetente: dict, edital_info: dict, payload: dict) -> bytes:
    """remetente/edital_info/payload: mesmo formato de gerar_pdf_proposta
    (proposta_pdf.py) -- payload é o dict de _proposta_payload (main.py)."""
    doc = Document()
    secao = doc.sections[0]
    secao.left_margin = secao.right_margin = Cm(1.7)
    secao.top_margin = secao.bottom_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    emp = remetente.get("empresa") or {}
    end = remetente.get("endereco") or {}

    # ---- cabeçalho: logo (se houver) + nome/CNPJ ----
    logo = _decodificar_logo(remetente.get("logo_base64"))
    if logo:
        try:
            doc.add_picture(logo, width=Cm(2.2))
        except Exception:
            pass
    p = doc.add_paragraph()
    _texto(p, remetente.get("nome") or "", negrito=True, tamanho=13)
    if remetente.get("documento"):
        p2 = doc.add_paragraph()
        _texto(p2, f"CNPJ/CPF: {remetente['documento']}", tamanho=9, cor=_COR_MUTED)
    doc.add_paragraph().add_run().add_break()

    # ---- destinatário / identificação do edital ----
    p = doc.add_paragraph()
    _texto(p, edital_info.get("orgao") or "Órgão não informado", negrito=True, tamanho=11)
    linha_edital = " Nº: ".join(v for v in [edital_info.get("modalidade"), edital_info.get("id_externo")] if v)
    if linha_edital:
        doc.add_paragraph(linha_edital)

    # ---- identificação do proponente ----
    p = doc.add_paragraph()
    _texto(p, "IDENTIFICAÇÃO DO PROPONENTE:", negrito=True, tamanho=10)
    endereco_txt = ", ".join(v for v in [end.get("logradouro"), end.get("numero"), end.get("bairro")] if v)
    cidade_txt = " - ".join(v for v in [end.get("cidade"), end.get("uf")] if v)
    partes_id = [remetente.get("nome") or ""]
    if remetente.get("documento"):
        partes_id.append(f"inscrita no CNPJ/CPF sob o nº {remetente['documento']}")
    if endereco_txt or cidade_txt:
        sede = ", ".join(v for v in [endereco_txt, cidade_txt, end.get("cep")] if v)
        partes_id.append(f"com sede em {sede}")
    doc.add_paragraph(", ".join(partes_id) + ".")

    if emp.get("representante_legal"):
        p = doc.add_paragraph()
        _texto(p, "REPRESENTANTE LEGAL: ", negrito=True)
        rg = f", inscrito no RG nº {emp['representante_rg']}" if emp.get("representante_rg") else ""
        _texto(p, f"{emp['representante_legal']}{rg}.")

    doc.add_paragraph("Proposta de preços, conforme Termo de Referência do Edital em epígrafe, "
                      "nas seguintes condições:")

    # ---- tabela de itens (mesma ordem de colunas do PDF) ----
    itens = payload.get("itens") or []
    colunas = ["Nº", "Descrição", "UND", "Qtd.", "Valor unit.", "Valor total", "Fabricante", "Marca", "Modelo"]
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.style = "Table Grid"
    for i, titulo in enumerate(colunas):
        celula = tabela.rows[0].cells[i]
        celula.paragraphs[0].add_run(titulo).bold = True
        _celula_sombreada(celula)
    for it in itens:
        qtd = it.get("quantidade") or 0
        preco = it.get("preco_unit") or 0
        numero = it.get("numero")
        linha = tabela.add_row().cells
        valores = [
            str(numero) if numero is not None else "-",
            str(it.get("descricao") or ""),
            _abreviar_unidade(it.get("unidade_medida") or ""),
            f"{qtd:g}", _fmt_moeda(preco), _fmt_moeda(preco * qtd),
            str(it.get("fabricante") or ""), str(it.get("marca") or ""), str(it.get("modelo") or ""),
        ]
        for i, v in enumerate(valores):
            linha[i].text = v

    doc.add_paragraph()
    p = doc.add_paragraph()
    _texto(p, f"VALOR TOTAL: {_fmt_moeda(payload.get('total_venda') or 0)}", negrito=True, tamanho=11)

    # ---- condições padrão ----
    for rotulo in ("VALIDADE DA PROPOSTA COMERCIAL", "PRAZO DE ENTREGA", "PRAZO DE GARANTIA"):
        p = doc.add_paragraph()
        _texto(p, f"{rotulo}: ", negrito=True)
        _texto(p, "conforme condições do edital.")

    dados_banco = [v for v in [
        f"Banco {emp['banco_nome']}" if emp.get("banco_nome") else None,
        f"conta {emp['banco_conta']}" if emp.get("banco_conta") else None,
        f"agência {emp['banco_agencia']}" if emp.get("banco_agencia") else None,
    ] if v]
    if dados_banco:
        p = doc.add_paragraph()
        _texto(p, "DADOS BANCÁRIOS: ", negrito=True)
        titular = f" ({remetente['nome']})" if remetente.get("nome") else ""
        _texto(p, " --- ".join(dados_banco) + titular + ".")

    observacoes = payload.get("observacoes")
    if observacoes:
        p = doc.add_paragraph()
        _texto(p, "Observações", negrito=True)
        p2 = doc.add_paragraph()
        _texto(p2, observacoes, cor=_COR_MUTED, tamanho=9.5)

    # ---- declaração ----
    p = doc.add_paragraph()
    _texto(p, "DECLARO ESTAR CIENTE E DE ACORDO COM O EDITAL E SEUS ANEXOS.", negrito=True)
    p2 = doc.add_paragraph()
    _texto(p2, _DECLARACAO, tamanho=8.5, cor=_COR_MUTED)

    # ---- local/data + assinatura ----
    hoje = datetime.date.today()
    cidade_data = ", ".join(v for v in [end.get("cidade"), end.get("uf")] if v)
    linha_data = f"{cidade_data + ', ' if cidade_data else ''}{hoje.day} de {_MESES[hoje.month-1]} de {hoje.year}."
    p = doc.add_paragraph(linha_data)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph("_" * 45)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _texto(p, remetente.get("nome") or "", negrito=True)
    if remetente.get("documento"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _texto(p, f"CNPJ/CPF nº {remetente['documento']}", tamanho=8.5)
    if emp.get("representante_legal"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rg = f" - RG: {emp['representante_rg']}" if emp.get("representante_rg") else ""
        _texto(p, f"{emp['representante_legal']}{rg}", tamanho=8.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
